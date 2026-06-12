"""Local document text extraction and OCR fallback."""
from __future__ import annotations

import html
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, List, Optional, Sequence

from service.mineru_document_parser import parse_with_mineru_hybrid


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".csv",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".py",
    ".js",
    ".ts",
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".pptx"}

REPO_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_VL_MODEL_DIR = REPO_ROOT / "models" / "PaddleOCR-VL-1.6"

_PADDLEOCR_VL_PIPELINE = None


class LocalDocumentParseError(RuntimeError):
    """Raised when local document parsing cannot produce text."""


@dataclass
class ParseResult:
    text: str
    method: str
    file_type: str
    page_count: int = 0
    ocr_pages: int = 0
    warnings: List[str] = field(default_factory=list)


def parse_document(file_path: str, file_name: Optional[str] = None) -> ParseResult:
    """Extract searchable text using direct parsers first, OCR only when needed."""
    path = Path(file_path)
    if not path.exists():
        raise LocalDocumentParseError(f"File does not exist: {file_path}")

    name = file_name or path.name
    ext = Path(name).suffix.lower() or path.suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(path, ext)
    if ext == ".docx":
        return _extract_docx(path, ext)
    if ext == ".xlsx":
        return _extract_xlsx(path, ext)
    if ext == ".pptx":
        return _extract_pptx(path, ext)
    if ext in TEXT_EXTENSIONS:
        return _extract_text_file(path, ext)
    if ext in IMAGE_EXTENSIONS:
        warnings: List[str] = []
        text = _ocr_image_paths([path], warnings)
        return ParseResult(
            text=_normalize_text(text),
            method="image:mineru_hybrid_medium",
            file_type=ext,
            page_count=1,
            ocr_pages=1,
            warnings=warnings,
        )

    if ext in {".doc", ".xls", ".ppt"}:
        raise LocalDocumentParseError(
            f"Legacy binary format {ext} is not supported by the local parser; "
            "convert it to docx/xlsx/pptx first."
        )
    raise LocalDocumentParseError(f"Unsupported file type: {ext or 'unknown'}")


def _extract_pdf(path: Path, ext: str) -> ParseResult:
    warnings: List[str] = []
    page_count = _count_pdf_pages(path, warnings)

    text = _extract_pdf_with_pdftotext(path, warnings)
    if _has_meaningful_text(text):
        return ParseResult(
            text=_normalize_text(text),
            method="pdf:pdftotext",
            file_type=ext,
            page_count=page_count,
            warnings=warnings,
        )

    text = _extract_pdf_with_pdfplumber(path, warnings)
    if _has_meaningful_text(text):
        return ParseResult(
            text=_normalize_text(text),
            method="pdf:pdfplumber",
            file_type=ext,
            page_count=page_count,
            warnings=warnings,
        )

    text, ocr_pages = _extract_pdf_with_ocr(path, warnings)
    if not _has_meaningful_text(text):
        raise LocalDocumentParseError("PDF text extraction and OCR returned empty content.")
    return ParseResult(
        text=_normalize_text(text),
        method="pdf:ocr:mineru_hybrid_medium",
        file_type=ext,
        page_count=page_count or ocr_pages,
        ocr_pages=ocr_pages,
        warnings=warnings,
    )


def _extract_pdf_with_pdftotext(path: Path, warnings: List[str]) -> str:
    if not shutil.which("pdftotext"):
        warnings.append("pdftotext is not installed; skipped direct PDF extraction.")
        return ""

    timeout = _env_int("PDFTOTEXT_TIMEOUT", 120)
    try:
        completed = subprocess.run(
            ["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
            check=False,
            capture_output=True,
            timeout=timeout,
        )
    except Exception as exc:
        warnings.append(f"pdftotext failed: {exc}")
        return ""

    if completed.returncode != 0:
        error = completed.stderr.decode("utf-8", errors="ignore").strip()
        warnings.append(f"pdftotext returned {completed.returncode}: {error[:200]}")
        return ""
    return completed.stdout.decode("utf-8", errors="ignore")


def _extract_pdf_with_pdfplumber(path: Path, warnings: List[str]) -> str:
    try:
        import pdfplumber
    except Exception as exc:
        warnings.append(f"pdfplumber is unavailable: {exc}")
        return ""

    try:
        page_texts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                page_texts.append(page.extract_text() or "")
        return "\n\n".join(page_texts)
    except Exception as exc:
        warnings.append(f"pdfplumber failed: {exc}")
        return ""


def _extract_pdf_with_ocr(path: Path, warnings: List[str]) -> tuple[str, int]:
    backend = os.getenv("LOCAL_OCR_BACKEND", "mineru_hybrid_medium").strip().lower()
    if backend in {"", "mineru", "mineru_hybrid", "mineru_hybrid_medium", "hybrid_medium"}:
        return _ocr_pdf_with_mineru_hybrid(path, warnings)

    if not shutil.which("pdftoppm"):
        raise LocalDocumentParseError(
            "Scanned PDF requires pdftoppm to render pages before OCR."
        )

    dpi = _env_int("LOCAL_OCR_DPI", 200)
    max_pages = _env_int("LOCAL_OCR_MAX_PAGES", 0)
    with tempfile.TemporaryDirectory(prefix="local_pdf_ocr_") as tmp_dir:
        prefix = Path(tmp_dir) / "page"
        cmd = ["pdftoppm", "-png", "-r", str(dpi)]
        if max_pages > 0:
            cmd.extend(["-f", "1", "-l", str(max_pages)])
        cmd.extend([str(path), str(prefix)])

        completed = subprocess.run(cmd, check=False, capture_output=True)
        if completed.returncode != 0:
            error = completed.stderr.decode("utf-8", errors="ignore").strip()
            raise LocalDocumentParseError(f"pdftoppm failed: {error[:300]}")

        image_paths = sorted(Path(tmp_dir).glob("page-*.png"))
        if not image_paths:
            raise LocalDocumentParseError("pdftoppm did not render any pages.")
        if max_pages > 0:
            warnings.append(f"OCR limited to first {max_pages} PDF pages.")
        return _ocr_image_paths(image_paths, warnings), len(image_paths)


def _ocr_pdf_with_mineru_hybrid(path: Path, warnings: List[str]) -> tuple[str, int]:
    max_pages = _env_int("LOCAL_OCR_MAX_PAGES", 0)
    with tempfile.TemporaryDirectory(prefix="local_pdf_mineru_") as tmp_dir:
        try:
            result = parse_with_mineru_hybrid(path, output_dir=Path(tmp_dir) / "mineru_hybrid")
        except Exception as exc:
            raise LocalDocumentParseError(f"MinerU hybrid OCR failed for PDF {path.name}: {exc}") from exc
        pages = result.pages[:max_pages] if max_pages > 0 else result.pages
        if max_pages > 0 and len(result.pages) > max_pages:
            warnings.append(f"OCR output limited to first {max_pages} PDF pages.")
        text = "\n\n".join((page.markdown.strip() or page.text.strip()) for page in pages if (page.markdown.strip() or page.text.strip()))
        return text, len(pages)


def _count_pdf_pages(path: Path, warnings: List[str]) -> int:
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            return len(pdf.pages)
    except Exception as exc:
        warnings.append(f"Could not count PDF pages: {exc}")
        return 0


def _extract_docx(path: Path, ext: str) -> ParseResult:
    try:
        from docx import Document as DocxDocument
    except Exception as exc:
        raise LocalDocumentParseError(f"python-docx is unavailable: {exc}") from exc

    doc = DocxDocument(str(path))
    parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [_normalize_inline_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))

    text = "\n\n".join(parts)
    if not _has_meaningful_text(text):
        raise LocalDocumentParseError("DOCX did not contain extractable text.")
    return ParseResult(text=_normalize_text(text), method="docx:python-docx", file_type=ext)


def _extract_xlsx(path: Path, ext: str) -> ParseResult:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise LocalDocumentParseError(f"openpyxl is unavailable: {exc}") from exc

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: List[str] = []
    try:
        for ws in wb.worksheets:
            sheet_lines = [f"# {ws.title}"]
            for row in ws.iter_rows(values_only=True):
                values = [_cell_to_text(value) for value in row]
                values = [value for value in values if value]
                if values:
                    sheet_lines.append(" | ".join(values))
            if len(sheet_lines) > 1:
                parts.append("\n".join(sheet_lines))
    finally:
        wb.close()

    text = "\n\n".join(parts)
    if not _has_meaningful_text(text):
        raise LocalDocumentParseError("XLSX did not contain extractable text.")
    return ParseResult(text=_normalize_text(text), method="xlsx:openpyxl", file_type=ext)


def _extract_pptx(path: Path, ext: str) -> ParseResult:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise LocalDocumentParseError(f"python-pptx is unavailable: {exc}") from exc

    prs = Presentation(str(path))
    parts: List[str] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_parts = [f"# Slide {slide_idx}"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = _normalize_inline_text(shape.text)
                if text:
                    slide_parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [_normalize_inline_text(cell.text) for cell in row.cells]
                    if any(cells):
                        slide_parts.append(" | ".join(cells))
        if len(slide_parts) > 1:
            parts.append("\n".join(slide_parts))

    text = "\n\n".join(parts)
    if not _has_meaningful_text(text):
        raise LocalDocumentParseError("PPTX did not contain extractable text.")
    return ParseResult(text=_normalize_text(text), method="pptx:python-pptx", file_type=ext)


def _extract_text_file(path: Path, ext: str) -> ParseResult:
    raw_text = _read_text(path)
    if ext in {".html", ".htm"}:
        raw_text = _html_to_text(raw_text)
    text = _normalize_text(raw_text)
    if not _has_meaningful_text(text):
        raise LocalDocumentParseError("Text file is empty after decoding.")
    return ParseResult(text=text, method=f"{ext.lstrip('.')}:direct-read", file_type=ext)


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    try:
        import chardet

        detected = chardet.detect(data).get("encoding")
        if detected:
            return data.decode(detected, errors="ignore")
    except Exception:
        pass

    return data.decode("utf-8", errors="ignore")


def _html_to_text(value: str) -> str:
    try:
        from html_text import extract_text

        extracted = extract_text(value)
        if extracted and extracted.strip():
            return extracted
    except Exception:
        pass

    without_scripts = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", value)
    without_tags = re.sub(r"(?s)<[^>]+>", " ", without_scripts)
    return html.unescape(without_tags)


def _ocr_image_paths(image_paths: Sequence[Path], warnings: List[str]) -> str:
    backend = os.getenv("LOCAL_OCR_BACKEND", "mineru_hybrid_medium").strip().lower()
    if backend in {"", "mineru", "mineru_hybrid", "mineru_hybrid_medium", "hybrid_medium"}:
        return _ocr_with_mineru_hybrid(image_paths, warnings)
    if backend in {"paddleocr_vl", "paddle_vl", "vl"}:
        return _ocr_with_paddleocr_vl(image_paths)
    if backend in {"paddleocr", "paddle"}:
        return _ocr_with_classic_paddleocr(image_paths)
    raise LocalDocumentParseError(f"Unsupported LOCAL_OCR_BACKEND: {backend}")


def _ocr_with_mineru_hybrid(image_paths: Sequence[Path], warnings: List[str]) -> str:
    page_texts: List[str] = []
    with tempfile.TemporaryDirectory(prefix="local_mineru_ocr_") as tmp_dir:
        tmp_root = Path(tmp_dir)
        for index, image_path in enumerate(image_paths, start=1):
            try:
                result = parse_with_mineru_hybrid(image_path, output_dir=tmp_root / f"page_{index}")
            except Exception as exc:
                raise LocalDocumentParseError(f"MinerU hybrid OCR failed for {image_path.name}: {exc}") from exc
            for page in result.pages:
                text = page.markdown.strip() or page.text.strip()
                if text:
                    page_texts.append(text)
    if not page_texts:
        warnings.append("MinerU hybrid OCR returned empty text.")
    return "\n\n".join(page_texts)


def _ocr_with_paddleocr_vl(image_paths: Sequence[Path]) -> str:
    pipeline = _get_paddleocr_vl_pipeline()
    use_layout_detection = _env_bool("PADDLEOCR_USE_LAYOUT_DETECTION", False)
    use_queues = _env_bool("PADDLEOCR_USE_QUEUES", False)
    max_new_tokens = _env_optional_int("PADDLEOCR_VL_MAX_NEW_TOKENS")

    kwargs: dict[str, Any] = {
        "use_layout_detection": use_layout_detection,
        "prompt_label": "ocr",
        "use_queues": use_queues,
    }
    if max_new_tokens:
        kwargs["max_new_tokens"] = max_new_tokens

    results = pipeline.predict([str(path) for path in image_paths], **kwargs)
    page_texts = [_extract_paddleocr_vl_result_text(result) for result in results]
    return "\n\n".join(text for text in page_texts if text.strip())


def _get_paddleocr_vl_pipeline():
    global _PADDLEOCR_VL_PIPELINE
    if _PADDLEOCR_VL_PIPELINE is not None:
        return _PADDLEOCR_VL_PIPELINE

    model_dir = Path(os.getenv("PADDLEOCR_VL_MODEL_DIR", str(DEFAULT_VL_MODEL_DIR))).expanduser()
    if not model_dir.exists():
        raise LocalDocumentParseError(
            f"PaddleOCR-VL model directory not found: {model_dir}. "
            "Set PADDLEOCR_VL_MODEL_DIR to the local PaddleOCR-VL-1.6 model path."
        )

    try:
        from paddleocr import PaddleOCRVL
    except Exception as exc:
        raise LocalDocumentParseError(f"paddleocr is unavailable: {exc}") from exc

    _PADDLEOCR_VL_PIPELINE = PaddleOCRVL(
        pipeline_version=os.getenv("PADDLEOCR_VL_VERSION", "v1.6"),
        vl_rec_model_dir=str(model_dir),
        vl_rec_backend=os.getenv("PADDLEOCR_VL_BACKEND", "native"),
        use_doc_orientation_classify=False,
        use_doc_unwarping=False,
        use_layout_detection=_env_bool("PADDLEOCR_USE_LAYOUT_DETECTION", False),
        use_chart_recognition=False,
        use_seal_recognition=False,
        use_ocr_for_image_block=True,
        format_block_content=False,
        merge_layout_blocks=True,
        use_queues=_env_bool("PADDLEOCR_USE_QUEUES", False),
    )
    return _PADDLEOCR_VL_PIPELINE


def _extract_paddleocr_vl_result_text(result: Any) -> str:
    try:
        markdown = result.markdown
        if isinstance(markdown, dict):
            text = markdown.get("markdown_texts")
            if isinstance(text, str) and text.strip():
                return text
    except Exception:
        pass

    try:
        data = result.json
        text = "\n".join(_iter_block_content(data))
        if text.strip():
            return text
    except Exception:
        pass

    if isinstance(result, dict):
        text = "\n".join(_iter_block_content(result))
        if text.strip():
            return text
    return str(result)


def _iter_block_content(value: Any) -> Iterable[str]:
    if isinstance(value, dict):
        content = value.get("block_content")
        if isinstance(content, str) and content.strip():
            yield content.strip()
        for key in ("res", "parsing_res_list", "pages", "results"):
            if key in value:
                yield from _iter_block_content(value[key])
    elif isinstance(value, list):
        for item in value:
            yield from _iter_block_content(item)


def _ocr_with_classic_paddleocr(image_paths: Sequence[Path]) -> str:
    try:
        from paddleocr import PaddleOCR
    except Exception as exc:
        raise LocalDocumentParseError(f"paddleocr is unavailable: {exc}") from exc

    ocr = PaddleOCR(lang=os.getenv("PADDLEOCR_LANG", "ch"))
    texts: List[str] = []
    for image_path in image_paths:
        output = ocr.predict(str(image_path))
        texts.extend(_iter_classic_ocr_text(output))
    return "\n".join(texts)


def _iter_classic_ocr_text(value: Any) -> Iterable[str]:
    if isinstance(value, dict):
        for key in ("rec_texts", "text", "label"):
            item = value.get(key)
            if isinstance(item, str) and item.strip():
                yield item.strip()
            elif isinstance(item, list):
                for text in item:
                    if isinstance(text, str) and text.strip():
                        yield text.strip()
        for item in value.values():
            if isinstance(item, (dict, list, tuple)):
                yield from _iter_classic_ocr_text(item)
    elif isinstance(value, (list, tuple)):
        if len(value) >= 2 and isinstance(value[1], (list, tuple)) and value[1]:
            candidate = value[1][0]
            if isinstance(candidate, str) and candidate.strip():
                yield candidate.strip()
        for item in value:
            if isinstance(item, (dict, list, tuple)):
                yield from _iter_classic_ocr_text(item)


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _normalize_inline_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _normalize_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _has_meaningful_text(value: str) -> bool:
    compact = re.sub(r"\s+", "", value or "")
    min_chars = _env_int("LOCAL_TEXT_MIN_CHARS", 20)
    if len(compact) < min_chars:
        return False
    useful_chars = re.findall(r"[\w\u4e00-\u9fff]", compact)
    return len(useful_chars) >= min_chars


def _env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


def _env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if value is None:
        return default
    try:
        return int(value)
    except ValueError:
        return default


def _env_optional_int(name: str) -> Optional[int]:
    value = os.getenv(name)
    if not value:
        return None
    try:
        return int(value)
    except ValueError:
        return None
